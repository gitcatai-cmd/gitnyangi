# K01 댓글 자료 — 이력서 서식 워드 파일 3종(기본·경력직·아르바이트)을 만든다, 칸마다 (괄호) 안내만 채워 둔다
import pathlib
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
OUT = pathlib.Path(__file__).parent
BLUE = RGBColor(0x2B, 0x47, 0xB5); GRAY = RGBColor(0x66, 0x6E, 0x7A)
W = 17.0  # 쓸 수 있는 폭(cm)

def base():
    d = Document(); s = d.styles['Normal']; s.font.name = '맑은 고딕'; s.font.size = Pt(10.5)
    s.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    for sec in d.sections: sec.top_margin = sec.bottom_margin = Cm(1.8); sec.left_margin = sec.right_margin = Cm(2.0)
    return d
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexc); tcPr.append(sh)
def table(d, rows, cols, widths):
    t = d.add_table(rows=rows, cols=cols); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for row in t.rows:
        for j, c in enumerate(row.cells): c.width = Cm(widths[j])
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[j] * 567)))  # 기준 칸 너비도 맞춘다(안 맞추면 반반으로 그리는 프로그램이 있다)
    return t
def height(row, cm):
    trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(int(cm * 567))); h.set(qn('w:hRule'), 'atLeast'); trPr.append(h)
def title(d, t, sub):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(24); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(); r = p.add_run(sub); r.font.size = Pt(9.5); r.font.color.rgb = GRAY; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
def head(d, t):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(12); r = p.add_run(t); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = BLUE
def kv(d, rows):
    t = table(d, len(rows), 2, (3.6, W - 3.6))
    for i, (a, b) in enumerate(rows):
        t.cell(i, 0).text = a; t.cell(i, 1).text = b; shade(t.cell(i, 0), 'EEF2FF')
        t.cell(i, 0).paragraphs[0].runs[0].bold = True; height(t.rows[i], 0.8)
def grid(d, hdr, n, widths=None):
    widths = widths or [W / len(hdr)] * len(hdr)
    t = table(d, n + 1, len(hdr), widths)
    for j, h in enumerate(hdr):
        t.cell(0, j).text = h; shade(t.cell(0, j), 'EEF2FF'); t.cell(0, j).paragraphs[0].runs[0].bold = True
        t.cell(0, j).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i in range(1, n + 1): height(t.rows[i], 0.9)
def note(d, t):
    p = d.add_paragraph(); r = p.add_run(t); r.font.size = Pt(9); r.font.color.rgb = GRAY
def foot(d):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(18)
    r = p.add_run('(괄호) 안의 글자를 지우고 내 내용을 쓰세요. 필요 없는 칸은 줄째 지워도 돼요. · 깃냥이 무료 자료 · CC BY 4.0'); r.font.size = Pt(8.5); r.font.color.rgb = GRAY

P = [('이름', '(이름)'), ('생년월일', '(예: 1980년 1월 1일)'), ('연락처', '(휴대전화 번호)'), ('이메일', '(이메일 주소)'), ('주소', '(사는 곳 — 시·구까지만 써도 돼요)')]
D3 = (4.0, 6.0, 7.0)
d = base(); title(d, '이 력 서', '기본 서식'); kv(d, P)
head(d, '학력'); grid(d, ['기간', '학교 이름', '전공 · 졸업 여부'], 3, D3)
head(d, '경력'); grid(d, ['기간', '회사 이름', '맡은 일'], 3, D3)
head(d, '자격증'); grid(d, ['딴 날', '자격증 이름', '발급 기관'], 3, D3)
head(d, '자기소개'); d.add_paragraph('(나를 소개하는 글 — 세 문장이면 충분해요. 무엇을 잘하는지, 어떤 사람인지, 왜 이 일을 하고 싶은지)')
d.add_paragraph(); p = d.add_paragraph('위 내용은 사실과 다름없습니다.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = d.add_paragraph('(날짜)          작성자 (이름)   (서명)'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
foot(d); d.save(OUT / '이력서_1_기본.docx')

d = base(); title(d, '이 력 서', '경력직 서식'); kv(d, P[:4])
head(d, '경력 요약'); d.add_paragraph('(한 줄로 — 예) 사무·총무 15년, 거래처 관리와 문서 작업에 강합니다)')
head(d, '경력 사항')
for i in range(2):
    kv(d, [('기간', '(예: 2015.03 ~ 2024.12)'), ('회사 · 부서', '(회사 이름 · 부서)'), ('직급', '(직급)'), ('한 일', '(맡은 일 두세 가지)'), ('성과', '(숫자가 있으면 더 좋아요 — 예) 거래처 30곳 관리)')]); d.add_paragraph()
head(d, '자격증 · 교육'); grid(d, ['딴 날', '이름', '발급 기관'], 3, D3)
head(d, '자기소개'); d.add_paragraph('(경력에서 배운 것과, 새 회사에서 할 수 있는 일)')
foot(d); d.save(OUT / '이력서_2_경력직.docx')

d = base(); title(d, '이 력 서', '아르바이트 서식'); kv(d, P[:3] + [('나이', '(만 나이)')])
head(d, '일할 수 있는 때'); grid(d, ['월', '화', '수', '목', '금', '토', '일'], 1)
note(d, '(칸에 ○ 표시하고, 가능한 시간도 적어요 — 예) 오전 9시~오후 3시)')
head(d, '일해 본 곳'); grid(d, ['기간', '가게 · 회사', '한 일'], 3, D3)
head(d, '한마디'); d.add_paragraph('(성실함, 일할 수 있는 날, 집에서 가깝다는 것 — 짧게)')
foot(d); d.save(OUT / '이력서_3_아르바이트.docx')
print('ok')
